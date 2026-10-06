"""Load files other than CSV: Excel, SPSS, Stata and Parquet.

Each reader builds a raw DuckDB table from the file, then `finish_load` gives the columns safe
names and re-types text that is really numbers or dates, as for a CSV.

- **Excel (.xlsx):** the first sheet with data; its first non-empty row is the header. Other
  sheets are listed in a warning. Formula cells use their last calculated value.
- **SPSS (.sav, .zsav, .por) and Stata (.dta):** variable labels become column descriptions
  (shown as coming from the file, like a data dictionary). A variable whose values all have
  value labels is stored as its labels ("Female", not 2); one with only some codes labelled
  (a 1-5 scale with labelled ends) keeps its numbers and lists the labels in its description.
  User-defined missing values (SPSS "9 = Refused") become missing, and are reported.
- **Parquet:** read directly; nested columns are stored as text.
"""

import datetime as dt
import math
from collections.abc import Sequence
from pathlib import Path
from typing import Any

import duckdb
import numpy as np

from api.ingest.loader import LoadReport, finish_load, load_csv
from api.ingest.names import quote
from api.profiler.warnings import DataWarning, warning

CSV_SUFFIXES = (".csv", ".tsv", ".txt")
EXCEL_SUFFIXES = (".xlsx", ".xlsm")
SPSS_SUFFIXES = (".sav", ".zsav", ".por")
STATA_SUFFIXES = (".dta",)
PARQUET_SUFFIXES = (".parquet", ".pq")
SUFFIXES = CSV_SUFFIXES + EXCEL_SUFFIXES + SPSS_SUFFIXES + STATA_SUFFIXES + PARQUET_SUFFIXES
EXCEL_ERRORS = {"#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!", "#SPILL!"}
MAX_LABELS_IN_DESCRIPTION = 12


def file_kind(filename: str) -> str | None:
    """csv | excel | spss | stata | parquet, or None for an unsupported file."""
    name = filename.lower()
    for kind, suffixes in (
        ("csv", CSV_SUFFIXES),
        ("excel", EXCEL_SUFFIXES),
        ("spss", SPSS_SUFFIXES),
        ("stata", STATA_SUFFIXES),
        ("parquet", PARQUET_SUFFIXES),
    ):
        if name.endswith(suffixes):
            return kind
    return None


def load_file(con: duckdb.DuckDBPyConnection, path: Path, table_name: str, kind: str) -> LoadReport:
    if kind == "csv":
        return load_csv(con, path, table_name)
    raw = f"__raw_{table_name}"
    warnings: list[DataWarning] = []
    descriptions: dict[str, str] = {}
    if kind == "excel":
        columns = _read_excel(path, warnings)
        _create(con, raw, columns)
    elif kind in ("spss", "stata"):
        columns, descriptions = _read_stats(path, kind, warnings)
        _create(con, raw, columns)
    elif kind == "parquet":
        _read_parquet(con, path, raw)
    else:
        raise ValueError(f"Unsupported file type: {kind}")
    report = finish_load(con, raw, table_name, warnings)
    by_original = {c.original_name: c.name for c in report.columns}
    report.descriptions = {by_original[k]: v for k, v in descriptions.items() if k in by_original}
    return report


def _create(con: duckdb.DuckDBPyConnection, raw: str, columns: dict[str, list[Any]]) -> None:
    if not columns:
        raise ValueError("The file has no columns.")
    # DuckDB scans a dict of NumPy object arrays directly, keeping each column's type.
    arrays = {name: np.array(values, dtype=object) for name, values in columns.items()}  # noqa: F841
    con.execute(f"CREATE OR REPLACE TABLE {quote(raw)} AS SELECT * FROM arrays")


def _as_int_if_whole(values: Sequence[Any]) -> list[Any]:
    """Floats that are all whole numbers (codes, counts, years) become integers."""
    nums = [v for v in values if v is not None]
    if nums and all(
        isinstance(v, float) and math.isfinite(v) and v.is_integer() and abs(v) < 2**53
        for v in nums
    ):
        return [None if v is None else int(v) for v in values]
    return list(values)


# --- Excel -----------------------------------------------------------------------------------


def _read_excel(path: Path, warnings: list[DataWarning]) -> dict[str, list[Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = []
        for ws in wb.worksheets:
            rows = [
                r for r in ws.iter_rows(values_only=True) if any(_cell(v) is not None for v in r)
            ]
            if rows:
                sheets.append((ws.title, rows))
    finally:
        wb.close()
    if not sheets:
        raise ValueError("The workbook has no data.")
    title, rows = sheets[0]
    if len(sheets) > 1:
        others = ", ".join(f"'{t}'" for t, _ in sheets[1:])
        warnings.append(
            warning(
                "other_sheets",
                "info",
                f"Only the first sheet with data, '{title}', was loaded. Other sheets: {others}. "
                "Save each sheet as its own file to add it.",
            )
        )
    width = max(len(r) for r in rows)
    header = [_cell(v) for v in rows[0]] + [None] * (width - len(rows[0]))
    names = [str(h).strip() if h is not None else f"column{i}" for i, h in enumerate(header)]
    columns: dict[str, list[Any]] = {}
    for i, name in enumerate(names):
        while name in columns:
            name += "_"
        values = [_cell(r[i]) if i < len(r) else None for r in rows[1:]]
        columns[name] = _excel_column(values)
    return columns


def _cell(v: Any) -> Any:
    if isinstance(v, str):
        v = v.strip()
        return None if v == "" or v in EXCEL_ERRORS else v
    if isinstance(v, float) and not math.isfinite(v):
        return None
    if isinstance(v, dt.time):
        return v.isoformat()
    return v


def _excel_column(values: list[Any]) -> list[Any]:
    present = [v for v in values if v is not None]
    if not present:
        return [None] * len(values)
    if all(isinstance(v, dt.datetime) for v in present):
        if all(v.time() == dt.time(0) for v in present):
            return [v.date() if v is not None else None for v in values]
        return values
    if all(isinstance(v, bool) for v in present):
        return values
    if all(isinstance(v, int | float) and not isinstance(v, bool) for v in present):
        return _as_int_if_whole([None if v is None else float(v) for v in values])
    # Mixed types: text, re-typed later if nearly all values are numbers or dates.
    return [None if v is None else _text(v) for v in values]


def _text(v: Any) -> str:
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat(sep=" ")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


# --- SPSS and Stata --------------------------------------------------------------------------


def _read_stats(
    path: Path, kind: str, warnings: list[DataWarning]
) -> tuple[dict[str, list[Any]], dict[str, str]]:
    import pyreadstat

    if kind == "stata":
        data, meta = pyreadstat.read_dta(str(path), output_format="dict")
    elif path.suffix.lower() == ".por":
        data, meta = pyreadstat.read_por(str(path), output_format="dict")
    else:
        data, meta = pyreadstat.read_sav(str(path), output_format="dict", user_missing=True)

    labels = dict(zip(meta.column_names, meta.column_labels or [], strict=False))
    value_labels: dict[str, dict[Any, str]] = meta.variable_value_labels or {}
    missing_ranges: dict[str, list[dict[str, Any]]] = getattr(meta, "missing_ranges", {}) or {}

    columns: dict[str, list[Any]] = {}
    descriptions: dict[str, str] = {}
    recoded_missing: list[str] = []
    labelled: list[str] = []
    for name in meta.column_names:
        values = [_plain(v) for v in data[name]]
        ranges = missing_ranges.get(name, [])
        if ranges:
            n = 0
            for i, v in enumerate(values):
                if isinstance(v, int | float) and any(r["lo"] <= v <= r["hi"] for r in ranges):
                    values[i] = None
                    n += 1
            if n:
                vl = value_labels.get(name, {})
                shown = [
                    f"{_code(r['lo'])}" + (f" = {vl[r['lo']]}" if r["lo"] in vl else "")
                    for r in ranges
                    if r["lo"] == r["hi"]
                ]
                recoded_missing.append(f"'{name}' ({n}: {', '.join(shown) or 'a range'})")
        description = (labels.get(name) or "").strip()
        codes = value_labels.get(name)
        present = {v for v in values if v is not None}
        if codes and present and present <= set(codes):
            values = [None if v is None else codes[v] for v in values]
            labelled.append(name)
        else:
            values = _as_int_if_whole(values)
            if codes:
                listed = list(codes.items())[:MAX_LABELS_IN_DESCRIPTION]
                text = "; ".join(f"{_code(k)} = {v}" for k, v in listed)
                description = f"{description}. Codes: {text}" if description else f"Codes: {text}"
        columns[name] = values
        if description:
            descriptions[name] = description[:2000]

    source = "Stata" if kind == "stata" else "SPSS"
    if recoded_missing:
        warnings.append(
            warning(
                "user_missing",
                "info",
                f"Values defined as missing in the {source} file were set to missing: "
                + "; ".join(recoded_missing)
                + ".",
            )
        )
    if labelled:
        warnings.append(
            warning(
                "value_labels",
                "info",
                f"Codes were replaced by their {source} value labels in: "
                + ", ".join(f"'{n}'" for n in labelled)
                + ".",
            )
        )
    return columns, descriptions


def _plain(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, np.generic):
        v = v.item()
    if isinstance(v, float) and not math.isfinite(v):
        return None
    if isinstance(v, str) and not v.strip():
        return None
    return v


def _code(v: Any) -> str:
    return str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)


# --- Parquet ---------------------------------------------------------------------------------


def _read_parquet(con: duckdb.DuckDBPyConnection, path: Path, raw: str) -> None:
    source = f"read_parquet('{str(path).replace(chr(39), chr(39) * 2)}')"
    described = con.execute(f"DESCRIBE SELECT * FROM {source}").fetchall()
    select = []
    for name, typ, *_ in described:
        nested = any(t in typ for t in ("[]", "STRUCT", "MAP", "UNION")) or typ.startswith("ENUM")
        col = quote(name)
        select.append(f"CAST({col} AS VARCHAR) AS {col}" if nested else col)
    con.execute(f"CREATE OR REPLACE TABLE {quote(raw)} AS SELECT {', '.join(select)} FROM {source}")
