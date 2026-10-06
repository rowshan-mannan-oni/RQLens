"""Export a literature table: CSV, Excel, Markdown (with citations) and BibTeX.

Rejected values and cells still running are left empty. "Not found" is written out, and
values whose citations failed the check are marked "(unverified)".
"""

import csv
import io
import re
import unicodedata
from dataclasses import dataclass
from typing import Any

from api.review.templates import TemplateColumn

MEDIA = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "md": "text/markdown; charset=utf-8",
    "bib": "application/x-bibtex; charset=utf-8",
}


@dataclass
class TableData:
    name: str
    columns: list[TemplateColumn]
    papers: list[dict[str, Any]]  # PaperOut dumps, ready papers only
    cells: list[dict[str, Any]]  # CellOut dumps

    def cell(self, paper_id: int, key: str) -> dict[str, Any] | None:
        if not hasattr(self, "_index"):
            self._index = {(c["paper_id"], c["column_key"]): c for c in self.cells}
        return self._index.get((paper_id, key))


def slug(text: str) -> str:
    ascii_text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", ascii_text.lower()).strip("-") or "literature-table"


def value_text(cell: dict[str, Any] | None) -> str:
    if cell is None or cell.get("review") == "rejected":
        return ""
    status = cell.get("status")
    if status == "not_found":
        return "Not found"
    if status not in ("done", "unverified"):
        return ""
    value = cell.get("value_json")
    if isinstance(value, list):
        text = "; ".join(str(v) for v in value)
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = "" if value is None else str(value)
    return f"{text} (unverified)" if status == "unverified" and text else text


def citations(cell: dict[str, Any] | None) -> list[dict[str, Any]]:
    if cell is None or cell.get("review") == "rejected" or not value_text(cell):
        return []
    if cell.get("source") == "user":
        return []
    return list(cell.get("citations_json") or [])


def paper_name(p: dict[str, Any]) -> str:
    return str(p.get("title") or p["filename"])


def render(data: TableData, fmt: str) -> tuple[bytes, str]:
    if fmt == "csv":
        return to_csv(data).encode("utf-8-sig"), MEDIA["csv"]  # BOM so Excel reads UTF-8
    if fmt == "xlsx":
        return to_xlsx(data), MEDIA["xlsx"]
    if fmt == "md":
        return to_markdown(data).encode(), MEDIA["md"]
    if fmt == "bib":
        return to_bibtex(data.papers).encode(), MEDIA["bib"]
    raise ValueError(f"Unknown format {fmt}")


def _citation_lines(data: TableData, paper: dict[str, Any]) -> list[str]:
    lines = []
    for col in data.columns:
        for c in citations(data.cell(paper["id"], col.key)):
            mark = "" if c.get("verified") else " (unverified)"
            lines.append(f'{col.label} \u2013 p. {c.get("page")} [{c.get("label")}]: '
                         f'"{c.get("quote", "")}"{mark}')  # fmt: skip
    return lines


def to_csv(data: TableData) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["File", *(c.label for c in data.columns), "Citations"])
    for p in data.papers:
        w.writerow(
            [
                p["filename"],
                *(value_text(data.cell(p["id"], c.key)) for c in data.columns),
                "\n".join(_citation_lines(data, p)),
            ]
        )
    return buf.getvalue()


def to_xlsx(data: TableData) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Table"
    header = ["File", *(c.label for c in data.columns)]
    ws.append(header)
    for p in data.papers:
        ws.append([p["filename"], *(value_text(data.cell(p["id"], c.key)) for c in data.columns)])
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="EEF0FE")
    for cell in ws[1]:
        cell.font, cell.fill = bold, fill
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"  # keep the file and first column (usually the title) in view
    for i, _ in enumerate(header, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = 22 if i == 1 else 40

    cites = wb.create_sheet("Citations")
    cites.append(["File", "Paper", "Column", "Passage", "Page", "Quote", "Verified"])
    for cell in cites[1]:
        cell.font, cell.fill = bold, fill
    for p in data.papers:
        for col in data.columns:
            for c in citations(data.cell(p["id"], col.key)):
                cites.append([p["filename"], paper_name(p), col.label, c.get("label"),
                              c.get("page"), c.get("quote"), bool(c.get("verified"))])  # fmt: skip
    for letter, width in zip("ABCDEFG", (22, 40, 22, 10, 6, 80, 9), strict=True):
        cites.column_dimensions[letter].width = width
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _md(text: str) -> str:
    return text.replace("|", "\\|").replace("\n", " ")


def to_markdown(data: TableData) -> str:
    """A table with numbered citation markers, and the quotes listed under it."""
    lines = [f"# {data.name}", ""]
    lines.append("| Paper | " + " | ".join(_md(c.label) for c in data.columns) + " |")
    lines.append("|---" * (len(data.columns) + 1) + "|")
    notes: list[str] = []
    for p in data.papers:
        row = [_md(paper_name(p))]
        for col in data.columns:
            cell = data.cell(p["id"], col.key)
            text = _md(value_text(cell))
            marks = []
            for c in citations(cell):
                notes.append(
                    f"{len(notes) + 1}. {_md(paper_name(p))}, p. {c.get('page')} "
                    f'[{c.get("label")}]: "{_md(str(c.get("quote", "")))}"'
                    + ("" if c.get("verified") else " (unverified)")
                )
                marks.append(f"[{len(notes)}]")
            row.append(f"{text} {' '.join(marks)}".strip())
        lines.append("| " + " | ".join(row) + " |")
    if notes:
        lines += ["", "## Citations", "", *notes]
    return "\n".join(lines) + "\n"


PROCEEDINGS = re.compile(r"proceedings|conference|symposium|workshop", re.IGNORECASE)


def _bib_escape(text: str) -> str:
    return re.sub(r"([&%$#_{}])", r"\\\1", text)


def to_bibtex(papers: list[dict[str, Any]]) -> str:
    entries = []
    used: set[str] = set()
    for p in papers:
        authors = p.get("authors_json") or []
        surname = re.sub(r"[^A-Za-z]", "", authors[0].split()[-1]) if authors else "anon"
        word = next(
            (w for w in re.findall(r"[A-Za-z]+", str(p.get("title") or p["filename"]))
             if len(w) > 3),
            "paper",
        )  # fmt: skip
        key = base = f"{surname.lower()}{p.get('year') or ''}{word.lower()}"
        n = 2
        while key in used:
            key, n = f"{base}{chr(ord('a') + n - 2)}", n + 1
        used.add(key)
        venue = p.get("venue")
        proceedings = bool(venue and PROCEEDINGS.search(venue))
        fields = {
            "title": p.get("title") or p["filename"],
            "author": " and ".join(authors) if authors else None,
            "year": p.get("year"),
            ("booktitle" if proceedings else "journal"): venue,
            "doi": p.get("doi"),
        }
        body = ",\n".join(
            f"  {k} = {{{_bib_escape(str(v))}}}" for k, v in fields.items() if v not in (None, "")
        )
        kind = "inproceedings" if proceedings else "article"
        entries.append(f"@{kind}{{{key},\n{body}\n}}")
    return "\n\n".join(entries) + ("\n" if entries else "")
