import csv
import io

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from api.review.export import TableData, render, to_bibtex, value_text
from api.review.templates import BUILTIN, TemplateColumn, slug, unique_keys

COLUMNS = [
    TemplateColumn(label="Paper title", metadata="title"),
    TemplateColumn(label="Metrics", kind="list"),
    TemplateColumn(label="Limitations", kind="list"),
    TemplateColumn(label="Sample size", kind="number"),
]
PAPERS = [
    {"id": 1, "filename": "a.pdf", "title": "Graph | Networks", "authors_json": ["Maria Okafor"],
     "year": 2023, "venue": "IEEE Transactions on Software Engineering", "doi": "10.1/x_y"},
    {"id": 2, "filename": "b.pdf", "title": None, "authors_json": None, "year": None,
     "venue": "Proceedings of the 2022 Conference on X", "doi": None},
]  # fmt: skip
CITE = {"label": "P1-S3", "page": 1, "quote": "We report F1", "verified": True}


def cell(paper_id, key, status="done", value=None, cites=(), source="llm", review=None):
    return {"paper_id": paper_id, "column_key": key, "status": status, "value_json": value,
            "citations_json": list(cites), "source": source, "review": review}  # fmt: skip


CELLS = [
    cell(1, "paper_title", value="Graph | Networks", source="metadata"),
    cell(1, "metrics", value=["F1", "AUC"], cites=[CITE]),
    cell(
        1,
        "limitations",
        status="unverified",
        value=["Only Java"],
        cites=[{**CITE, "verified": False}],
    ),
    cell(1, "sample_size", value=18.0),
    cell(2, "metrics", status="not_found"),
    cell(2, "limitations", value=["Rejected"], cites=[CITE], review="rejected"),
    cell(2, "sample_size", value="mine", cites=[CITE], source="user"),
    cell(2, "paper_title", status="running"),
]
DATA = TableData("My review: SE", COLUMNS, PAPERS, CELLS)


def test_value_text_marks_status():
    assert value_text(CELLS[1]) == "F1; AUC"
    assert value_text(CELLS[2]) == "Only Java (unverified)"
    assert value_text(CELLS[3]) == "18"
    assert value_text(CELLS[4]) == "Not found"
    assert value_text(CELLS[5]) == ""  # rejected
    assert value_text(CELLS[7]) == ""  # still running


def test_csv_has_values_and_citations():
    body, media = render(DATA, "csv")
    assert media.startswith("text/csv")
    rows = list(csv.reader(io.StringIO(body.decode("utf-8-sig"))))
    assert rows[0] == ["File", "Paper title", "Metrics", "Limitations", "Sample size", "Citations"]
    assert rows[1][:5] == ["a.pdf", "Graph | Networks", "F1; AUC", "Only Java (unverified)", "18"]
    assert 'Metrics \u2013 p. 1 [P1-S3]: "We report F1"' in rows[1][5]
    assert "(unverified)" in rows[1][5]
    assert rows[2][1:5] == ["", "Not found", "", "mine"]
    assert rows[2][5] == ""  # rejected and user values carry no AI citations


def test_xlsx_has_table_and_citation_sheets():
    body, _ = render(DATA, "xlsx")
    wb = load_workbook(io.BytesIO(body))
    assert wb.sheetnames == ["Table", "Citations"]
    table = [[c.value for c in r] for r in wb["Table"].iter_rows()]
    assert table[1][2] == "F1; AUC"
    cites = [[c.value for c in r] for r in wb["Citations"].iter_rows()]
    assert cites[1] == ["a.pdf", "Graph | Networks", "Metrics", "P1-S3", 1, "We report F1", True]
    assert wb["Table"].freeze_panes == "C2"


def test_markdown_escapes_pipes_and_numbers_citations():
    md = render(DATA, "md")[0].decode()
    assert "| Graph \\| Networks |" in md
    assert "F1; AUC [1]" in md
    assert "## Citations" in md
    assert '1. Graph \\| Networks, p. 1 [P1-S3]: "We report F1"' in md


def test_bibtex_entries():
    bib = to_bibtex(PAPERS)
    assert "@article{okafor2023graph," in bib
    assert "journal = {IEEE Transactions on Software Engineering}" in bib
    assert "doi = {10.1/x\\_y}" in bib
    assert "@inproceedings{anonpaper," in bib
    assert "booktitle = {Proceedings of the 2022 Conference on X}" in bib


def test_template_columns_validate_and_get_unique_keys():
    with pytest.raises(ValidationError):
        TemplateColumn(label="Design", kind="category", options=["RCT"])
    cols = unique_keys([TemplateColumn(label="Metrics"), TemplateColumn(label="metrics!")])
    assert [c.key for c in cols] == ["metrics", "metrics_2"]
    assert slug("Dataset(s) / used") == "dataset_s_used"
    lit = BUILTIN["builtin:literature_review"]
    assert len(lit.columns) == 13
    assert [c.key for c in lit.columns if c.metadata] == ["paper_title", "authors", "year"]
